"""Export a run's results: a printable page (export.html) and an Excel workbook (results.xlsx).

export.html puts the ranking and every client's conversation brief on one page, with a
page break before each client when printed. results.xlsx has the same content as sheets
(Ranking, Talking points, Key figures, Evidence, Warnings) for sorting and filtering.
SSNs are masked in both, as everywhere else.
"""
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .privacy import mask_ssn
from .report import _products, esc, money, page, summary_rows

XLSX = "results.xlsx"
PRODUCTS = ("aum", "annuity", "life")

EXPORT_CSS = """
.ex-head { display:flex; justify-content:space-between; align-items:flex-end; gap:16px; flex-wrap:wrap; }
.ex-actions { display:flex; gap:10px; flex-wrap:wrap; }
.ex-table { background:var(--card); border:1px solid var(--line); border-radius:16px; overflow:hidden; margin-top:18px; }
.ex-table td { font-size:13.5px; }
.ex-client { background:var(--card); border:1px solid var(--line); border-radius:18px; padding:20px 22px; margin-top:22px; }
.ex-client h3 { margin:0; font-size:20px; letter-spacing:-.015em; }
.ex-verdicts { display:flex; gap:8px; flex-wrap:wrap; margin:12px 0 4px; }
.ex-client h4 { margin:18px 0 6px; font-size:12px; text-transform:uppercase; letter-spacing:.06em; color:var(--ink-3); }
.ex-client ol { margin:0; padding-left:20px; } .ex-client li { margin:3px 0; }
.ex-tp td:first-child { white-space:nowrap; font-weight:650; }
.ex-kv { display:grid; grid-template-columns:repeat(auto-fit,minmax(190px,1fr)); gap:4px 18px; font-size:13.5px; }
.ex-kv div { display:flex; justify-content:space-between; gap:10px; border-bottom:1px dashed var(--line); padding:4px 0; }
.ex-kv span { color:var(--ink-2); } .ex-kv b { font-variant-numeric:tabular-nums; }
@media print {
  :root { --bg:#fff; --ink:#000; --ink-2:#333; --ink-3:#666; --line:#ccc; --card:#fff; --chip:#f2f2f2;
          --good:#0f7a55; --good-bg:#e3f6ee; --mid:#8a5300; --mid-bg:#fff1d6; --low:#555; --low-bg:#eee; --bad:#b0262b; }
  .ex-client { border:0; padding:0; break-before:page; }
  .ex-table { border:0; }
  .ex-tp tr { break-inside:avoid; }
}
"""


# ---------------------------------------------------------------------
# printable page
# ---------------------------------------------------------------------
def _pill(p):
    return (f"<span class='pill t-{p['tone']}'><span class='dot'></span>{esc(p['name'])}: {esc(p['verdict'])}"
            f" &middot; {esc(p['score_label'])}</span>")


def _client_section(r):
    if r["status"] != "graded":
        return (f"<section class='ex-client'><h3>{esc(r['client'])}</h3>"
                f"<p class='skip'>Not graded: {esc(r['reason'])}</p></section>")
    b, prods = r["brief"], _products(r)
    names = " & ".join(r["names"]) or r["client"]
    meta = [f"Tax year {r['year']}", r["filing_status"]] + [f"{n}, {a}" for n, a in r["ages"]]
    out = [f"<section class='ex-client'><h3>{esc(names)}</h3>"
           f"<div class='meta'>{''.join(f'<span class=tag>{esc(m)}</span>' for m in meta)}</div>"
           f"<div class='ex-verdicts'>{''.join(_pill(prods[k]) for k in PRODUCTS)}</div>"]
    if b["opener"]:
        out.append(f"<h4>What stood out on their {r['year']} return</h4><ol>"
                   + "".join(f"<li>{esc(x)}</li>" for x in b["opener"]) + "</ol>")
    rows = [(prods[k]["name"], t) for k in b["order"] for t in prods[k]["points"]]
    if rows:
        out.append("<h4>Talking points</h4><div class='wrap'><table class='ex-tp'><tr><th>Product</th><th>Point</th>"
                   "<th>Why it matters</th><th>Ask</th></tr>")
        for name, t in rows:
            extra = "".join(f"<div class='muted'>From the return: {esc(c)}</div>" for c in t.get("context", []))
            extra += "".join(f"<div class='muted'>Client note: &ldquo;{esc(n)}&rdquo;</div>"
                             for n in t.get("client_notes", []))
            out.append(f"<tr><td>{esc(name)}</td><td>{esc(t['headline'])}{extra}</td><td>{esc(t['why'])}</td>"
                       f"<td>{esc(t['ask'])}</td></tr>")
        out.append("</table></div>")
    else:
        out.append("<p class='muted'>Nothing on this return points to a conversation right now.</p>")
    kv = [(s["label"], money(s["value"])) for s in b["snapshot"]]
    kv += [(c["label"], f"{c['value']} ({c['sub']})") for c in b.get("tax_picture", [])]
    if kv:
        out.append("<h4>Household snapshot</h4><div class='ex-kv'>"
                   + "".join(f"<div><span>{esc(k)}</span><b>{esc(v)}</b></div>" for k, v in kv) + "</div>")
    if r["warnings"]:
        out.append(f"<h4>Check before the meeting</h4><ul>" + "".join(f"<li>{esc(w)}</li>" for w in r["warnings"]) + "</ul>")
    out.append("</section>")
    return "".join(out)


def export_html(results, when=None):
    when = when or datetime.now()
    rows = summary_rows(results)
    graded = [r for r in rows if r["status"] == "graded"]
    out = [f"<style>{EXPORT_CSS}</style>",
           "<div class='ex-head'><div><h1>Tax Scan results</h1>"
           f"<div class='sub'>{len(graded)} client{'s' if len(graded) != 1 else ''} graded &middot; "
           f"{when.strftime('%B %d, %Y')}</div></div>"
           "<div class='ex-actions noprint'><button class='btn primary' onclick='window.print()'>Print / save as PDF</button>"
           f"<a class='btn' href='{XLSX}' download>Download Excel</a></div></div>",
           "<div class='ex-table wrap'><table><tr><th>#</th><th>Client</th><th>Investment mgmt</th><th>Annuity</th>"
           "<th>Life insurance</th><th>What stood out</th></tr>"]
    for row in rows:
        if row["status"] != "graded":
            out.append(f"<tr><td>&ndash;</td><td>{esc(row['client'])}</td><td colspan='4' class='skip'>"
                       f"{esc(row['status'])}</td></tr>")
            continue
        p = _products(row["_r"])
        cells = "".join(f"<td><b class='c-{p[k]['tone']}'>{esc(p[k]['verdict'])}</b><br>"
                        f"<span class='muted'>{esc(p[k]['score_label'])}</span></td>" for k in PRODUCTS)
        out.append(f"<tr><td>{row['rank']}</td><td><b>{esc(row['names'] or row['client'])}</b></td>{cells}"
                   f"<td>{'<br>'.join(esc(x) for x in row['_r']['brief']['opener'][:2])}</td></tr>")
    out.append("</table></div>")
    ordered = [row["_r"] for row in graded] + [r for r in results if r["status"] != "graded"]
    out += [_client_section(r) for r in ordered]
    out.append("<p class='foot'>Prepared on this computer by Tax Scan from the clients&rsquo; tax documents. "
               "SSNs are masked. A starting point for a conversation, not advice.</p>")
    return page("Export", "".join(out))


# ---------------------------------------------------------------------
# Excel workbook
# ---------------------------------------------------------------------
HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="5B4CF0")
MONEY = '"$"#,##0'


def _cell_value(v):
    return mask_ssn(v) if isinstance(v, str) else v


def _sheet(wb, title, headers, rows, widths, money_cols=(), wrap_cols=()):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append([_cell_value(v) for v in row])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, str):
                c.data_type = "s"            # text that starts with "=" stays text, never a formula
            if c.column in money_cols and isinstance(c.value, (int, float)):
                c.number_format = MONEY
            c.alignment = Alignment(vertical="top", wrap_text=c.column in wrap_cols)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def export_workbook(results):
    wb = Workbook()
    wb.remove(wb.active)
    rows = summary_rows(results)
    _sheet(wb, "Ranking",
           ["Rank", "Client", "Names", "Tax year", "Investment score", "Investment verdict", "Annuity grade",
            "Annuity points", "Annuity verdict", "Life grade", "Life points", "Life verdict", "What stood out",
            "Warnings", "Status"],
           [[r["rank"], r["client"], r["names"], r["year"], r["aum_score"], r["aum_verdict"], r["annuity_grade"],
             r["annuity_points"], r["annuity_verdict"], r["life_grade"], r["life_points"], r["life_verdict"],
             "\n".join(r["_r"]["brief"]["opener"]) if r["status"] == "graded" else "", r["warnings"], r["status"]]
            for r in rows],
           [7, 18, 26, 9, 11, 20, 9, 9, 20, 9, 9, 20, 60, 10, 18], wrap_cols={13})
    graded = [r["_r"] for r in rows if r["status"] == "graded"]

    tps = []
    for r in graded:
        prods = _products(r)
        for k in r["brief"]["order"]:
            p = prods[k]
            for i, t in enumerate(p["points"], 1):
                tps.append([r["client"], p["name"], p["verdict"], i, t["headline"], t["why"], t["ask"],
                            "\n".join(t.get("context", [])), "\n".join(t.get("client_notes", []))])
    _sheet(wb, "Talking points", ["Client", "Product", "Verdict", "#", "Point", "Why it matters", "Ask",
                                  "From the return", "Client notes"],
           tps, [18, 20, 18, 4, 45, 55, 45, 40, 40], wrap_cols={5, 6, 7, 8, 9})

    figs = [[r["client"], f["label"], r["year"], f["value"], r["prior_year"] or "", f.get("prior"), f["source"]]
            for r in graded for f in r["figures"]]
    _sheet(wb, "Key figures", ["Client", "Figure", "Tax year", "Amount", "Prior year", "Prior amount", "Source"],
           figs, [18, 32, 9, 14, 10, 14, 50], money_cols={4, 6}, wrap_cols={7})

    sigs = [[r["client"], _products(r)[s["product"]]["name"], s["label"], s["value"], s["detail"], s["points"],
             "; ".join(s["sources"]), s["id"]] for r in graded for s in r["signals"]]
    _sheet(wb, "Evidence", ["Client", "Product", "Rule", "Value", "Detail", "Points", "Sources", "Rule id"],
           sigs, [18, 20, 40, 14, 40, 8, 50, 28], wrap_cols={3, 5, 7})

    warns = [[r["client"], w] for r in graded for w in r["warnings"]]
    warns += [[r["client"], f"Not graded: {r['reason']}"] for r in results if r["status"] != "graded"]
    _sheet(wb, "Warnings", ["Client", "Warning"], warns, [18, 100], wrap_cols={2})
    return wb


def write_export(results, out_dir):
    out_dir = Path(out_dir)
    (out_dir / "export.html").write_text(export_html(results), encoding="utf-8")
    export_workbook(results).save(out_dir / XLSX)
