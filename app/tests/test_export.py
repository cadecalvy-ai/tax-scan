"""Export: the printable page and the Excel workbook cover every client, with SSNs masked."""
import openpyxl
import pytest

from src import paths
from src.grader import run

SHEETS = ["Ranking", "Talking points", "Key figures", "Evidence", "Warnings"]


@pytest.fixture(scope="module")
def out(tmp_path_factory):
    d = tmp_path_factory.mktemp("export")
    results = run(paths.SAMPLE_CLIENTS, d, cache_dir=str(paths.APP / ".pytest_cache" / "ocr"), workers=4)
    return d, results


def test_workbook_sheets_and_rows(out):
    d, results = out
    wb = openpyxl.load_workbook(d / "results.xlsx")
    assert wb.sheetnames == SHEETS
    clients = {r["client"] for r in results}
    ranked = [row[1] for row in wb["Ranking"].iter_rows(min_row=2, values_only=True)]
    assert set(ranked) == clients and len(ranked) == len(clients)
    assert wb["Talking points"].max_row > len(clients)
    assert wb["Key figures"]["D2"].number_format == '"$"#,##0'


def test_workbook_has_no_formulas_or_ssns(out):
    import re
    d, _ = out
    wb = openpyxl.load_workbook(d / "results.xlsx")
    for ws in wb:
        for row in ws.iter_rows(min_row=2):
            for c in row:
                assert c.data_type != "f", f"{ws.title}!{c.coordinate} became a formula"
                if isinstance(c.value, str):
                    assert not re.search(r"\b\d{3}-\d{2}-\d{4}\b", c.value)


def test_printable_page_lists_every_client(out):
    d, results = out
    html = (d / "export.html").read_text(encoding="utf-8")
    assert "window.print()" in html and "results.xlsx" in html
    assert html.count("class='ex-client'") == len(results)
    assert "export.html" in (d / "summary.html").read_text(encoding="utf-8")
